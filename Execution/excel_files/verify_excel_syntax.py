"""
verify_excel_syntax.py
-----------------------
Validates an Excel mapping file against the syntax defined in
excel_syntax_guide.md before parsing.

Usage:
    python verify_excel_syntax.py <path_to_xlsx>

Programmatic:
    from verify_excel_syntax import verify
    result = verify("analytics_dw.public.dim_vehicle_master.xlsx")
    if not result.passed:
        print(result.summary())
"""

import os
import re
import sys
from dataclasses import dataclass, field
from typing import List, Optional

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required.  pip install openpyxl")
    sys.exit(1)


# ---------------------------------------------------------------------------
# Result container
# ---------------------------------------------------------------------------
@dataclass
class Issue:
    severity: str          # "ERROR" or "WARNING"
    sheet: Optional[str]   # None for file-level issues
    row: Optional[int]     # 1-based, None if not row-specific
    rule: str              # e.g. "S1", "D3"
    message: str


@dataclass
class VerificationResult:
    file_path: str
    issues: List[Issue] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return not any(i.severity == "ERROR" for i in self.issues)

    def summary(self) -> str:
        errors = [i for i in self.issues if i.severity == "ERROR"]
        warnings = [i for i in self.issues if i.severity == "WARNING"]
        lines = [
            f"{'PASS' if self.passed else 'FAIL'}  {self.file_path}",
            f"  Errors: {len(errors)}  |  Warnings: {len(warnings)}",
        ]
        for i in self.issues:
            loc = ""
            if i.sheet:
                loc += f"[{i.sheet}]"
            if i.row:
                loc += f" Row {i.row}"
            lines.append(f"  {i.severity:7s} {i.rule:4s} {loc}: {i.message}")
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Normalisation helpers
# ---------------------------------------------------------------------------
_WS = re.compile(r"\s+")

def _norm(val) -> str:
    if val is None:
        return ""
    return str(val).strip().lower()

def _clean(val) -> str:
    if val is None:
        return ""
    return str(val).strip()

def _is_blank(val) -> bool:
    s = _clean(val)
    return s == "" or s.lower() == "none" or s == "\xa0"


# ---------------------------------------------------------------------------
# Constants for this project's Excel format
# ---------------------------------------------------------------------------

# Filename: <db>.<schema>.<table>.xlsx
FILENAME_PATTERN = re.compile(
    r'^(?P<db>[^.]+)\.(?P<schema>[^.]+)\.(?P<table>.+)\.xlsx$',
    re.IGNORECASE,
)

# Expected column headers (Row 2)
EXPECTED_HEADERS = {
    "src_db", "src_table", "src_col_name", "src_dtype",
    "transform_type", "tgt_col_name", "tgt_dtype",
}

# Valid transform types
VALID_TRANSFORM_TYPES = {"direct", "rename", "cast", "derived", "constant", "drop"}

# Sheets to skip
SKIP_SHEETS = {"_joins", "_parameters"}

# Global transforms marker
GT_MARKER = "_global_transforms"

# Recognized global transform operations
RECOGNIZED_GLOBAL_OPS = {
    "trim", "upper", "lower", "replace", "regex_replace", "strip_whitespace",
    "cast", "ltrim", "rtrim", "lpad", "rpad",
}

# Valid join types
VALID_JOIN_TYPES = {"LEFT", "INNER", "RIGHT", "FULL"}

# Required joins columns
JOINS_REQUIRED_COLS = {
    "sheet_name", "join_alias", "join_table", "join_type",
    "left_key", "right_key", "fetch_columns",
}

# Parameters sheet constants
PARAMETERS_SHEET = "_parameters"
PARAMETERS_REQUIRED_HEADERS = {"sr_no", "parameter_name", "parameter_value"}


# ---------------------------------------------------------------------------
# Core verifier
# ---------------------------------------------------------------------------
def verify(file_path: str) -> VerificationResult:
    result = VerificationResult(file_path=file_path)

    # S1: File must be .xlsx
    if not file_path.lower().endswith(".xlsx"):
        result.issues.append(Issue("ERROR", None, None, "S1", "File must be .xlsx"))
        return result

    if not os.path.isfile(file_path):
        result.issues.append(Issue("ERROR", None, None, "S1", f"File not found: {file_path}"))
        return result

    # S0: Filename pattern
    basename = os.path.basename(file_path)
    if not FILENAME_PATTERN.match(basename):
        result.issues.append(Issue("ERROR", None, None, "S0",
            f"Filename '{basename}' does not match pattern "
            f"'<db>.<schema>.<table>.xlsx'"))

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        result.issues.append(Issue("ERROR", None, None, "S1", f"Cannot open workbook: {e}"))
        return result

    # Identify server sheets (non-skip, non-_joins)
    server_sheets = [
        s for s in wb.sheetnames
        if s.lower() not in SKIP_SHEETS and not s.lower().startswith("retired")
    ]

    # S2: At least 1 server sheet
    if not server_sheets:
        result.issues.append(Issue("ERROR", None, None, "S2",
            "No server/mapping sheets found"))
        return result

    sheet_targets = {}

    for sname in server_sheets:
        ws = wb[sname]
        _verify_server_sheet(ws, sname, result, sheet_targets)

    # Verify _joins sheet if present
    if "_joins" in [s.lower() for s in wb.sheetnames]:
        jws_name = next(s for s in wb.sheetnames if s.lower() == "_joins")
        _verify_joins_sheet(wb[jws_name], server_sheets, result)

    # Verify _parameters sheet if present
    if "_parameters" in [s.lower() for s in wb.sheetnames]:
        pws_name = next(s for s in wb.sheetnames if s.lower() == "_parameters")
        _verify_parameters_sheet(wb[pws_name], result)

    # Cross-sheet validation (X1-X3)
    if len(sheet_targets) > 1:
        _verify_cross_sheets(sheet_targets, result)

    return result


def _find_header_row(ws, max_col: int) -> Optional[int]:
    """Find the row containing column headers (Row 2 typically)."""
    for r in range(1, min((ws.max_row or 0) + 1, 5)):
        vals = {_norm(ws.cell(r, c).value) for c in range(1, max_col + 1)}
        if "src_db" in vals or "src_table" in vals or "tgt_col_name" in vals:
            return r
    return None


def _verify_server_sheet(ws, sname: str, result: VerificationResult,
                          sheet_targets: dict):
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if max_row < 3:
        result.issues.append(Issue("ERROR", sname, None, "S3",
            "Sheet has fewer than 3 rows (need header rows + data)"))
        return

    # Find header row
    header_row = _find_header_row(ws, max_col)
    if header_row is None:
        result.issues.append(Issue("ERROR", sname, None, "S3",
            "Cannot find header row with expected column names (src_db, tgt_col_name, etc.)"))
        return

    # Build header map
    hdr = {}
    for c in range(1, max_col + 1):
        n = _norm(ws.cell(header_row, c).value)
        if n:
            hdr[n] = c

    # S4: Required headers
    missing_hdrs = EXPECTED_HEADERS - set(hdr.keys())
    if missing_hdrs:
        result.issues.append(Issue("ERROR", sname, header_row, "S4",
            f"Missing required headers: {missing_hdrs}"))

    def _get(row, name):
        c = hdr.get(name)
        return _clean(ws.cell(row, c).value) if c else ""

    # Parse data rows
    data_start = header_row + 1
    tgt_columns = {}
    seen_pk = False
    gt_start = None

    for r in range(data_start, max_row + 1):
        # Check for global transforms marker
        cell_a = _clean(ws.cell(r, 1).value)
        if cell_a.lower().strip() == GT_MARKER:
            gt_start = r
            break

        # Skip blank rows
        row_vals = [_clean(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        if all(_is_blank(v) for v in row_vals):
            continue

        tgt_col = _get(r, "tgt_col_name")
        tgt_dtype = _get(r, "tgt_dtype")
        transform_type = _get(r, "transform_type").lower()
        transform_rule = _get(r, "transform_rule")
        tgt_is_pk = _get(r, "tgt_is_pk").upper()
        src_db = _get(r, "src_db")
        src_table = _get(r, "src_table")
        src_col = _get(r, "src_col_name")

        # D1: tgt_col_name not blank
        if _is_blank(tgt_col) and transform_type != "drop":
            result.issues.append(Issue("ERROR", sname, r, "D1",
                "tgt_col_name is blank"))
            continue

        # D2: tgt_dtype not blank
        if _is_blank(tgt_dtype) and transform_type != "drop":
            result.issues.append(Issue("ERROR", sname, r, "D2",
                f"tgt_dtype is blank for column '{tgt_col}'"))

        # D3: No duplicate target columns
        if tgt_col:
            col_upper = tgt_col.upper()
            if col_upper in tgt_columns:
                result.issues.append(Issue("ERROR", sname, r, "D3",
                    f"Duplicate target column '{tgt_col}'"))
            else:
                tgt_columns[col_upper] = {"dtype": tgt_dtype, "pk": tgt_is_pk}

        # D4: transform_type is valid
        if transform_type and transform_type not in VALID_TRANSFORM_TYPES:
            result.issues.append(Issue("ERROR", sname, r, "D4",
                f"Invalid transform_type '{transform_type}' for '{tgt_col}'. "
                f"Valid: {VALID_TRANSFORM_TYPES}"))

        # D5: constant rows should have blank source fields
        if transform_type == "constant":
            if not _is_blank(src_db) or not _is_blank(src_table) or not _is_blank(src_col):
                result.issues.append(Issue("WARNING", sname, r, "D5",
                    f"Constant column '{tgt_col}' has non-blank source fields"))

        # D6: non-constant rows should have source fields
        if transform_type and transform_type not in ("constant", "drop"):
            if _is_blank(src_table) and _is_blank(src_col):
                result.issues.append(Issue("WARNING", sname, r, "D6",
                    f"Column '{tgt_col}' (type={transform_type}) has blank source fields"))

        # D7: cast/derived/constant need transform_rule
        if transform_type in ("cast", "derived", "constant"):
            if _is_blank(transform_rule):
                result.issues.append(Issue("WARNING", sname, r, "D7",
                    f"Column '{tgt_col}' (type={transform_type}) has blank transform_rule"))

        # PK tracking
        if tgt_is_pk == "Y":
            seen_pk = True

    # D8: at least 1 PK
    if not seen_pk:
        result.issues.append(Issue("WARNING", sname, None, "D8",
            "No PK column found (tgt_is_pk = Y)"))

    # Validate global transforms if present
    if gt_start is not None:
        _verify_global_transforms(ws, sname, gt_start, max_row, max_col, hdr, result)

    sheet_targets[sname] = tgt_columns


def _verify_global_transforms(ws, sname: str, start_row: int, max_row: int,
                                max_col: int, hdr: dict, result: VerificationResult):
    """G1, G2: Validate global transforms section."""
    gt_hdr_row = start_row + 1
    if gt_hdr_row > max_row:
        return

    # Build GT header map
    gt_hdr = {}
    for c in range(1, max_col + 1):
        n = _norm(ws.cell(gt_hdr_row, c).value)
        if n:
            gt_hdr[n] = c

    expected_order = 1
    for r in range(gt_hdr_row + 1, max_row + 1):
        row_vals = [_clean(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        if all(_is_blank(v) for v in row_vals):
            continue

        order_col = gt_hdr.get("order")
        op_col = gt_hdr.get("operation")

        # G1: Order sequential
        if order_col:
            order_val = _clean(ws.cell(r, order_col).value)
            try:
                order_int = int(order_val)
                if order_int != expected_order:
                    result.issues.append(Issue("WARNING", sname, r, "G1",
                        f"Global transform Order expected {expected_order}, got {order_int}"))
                expected_order = order_int + 1
            except (ValueError, TypeError):
                result.issues.append(Issue("WARNING", sname, r, "G1",
                    f"Global transform Order '{order_val}' is not an integer"))

        # G2: Operation recognized
        if op_col:
            op_val = _clean(ws.cell(r, op_col).value).lower()
            if op_val and op_val not in RECOGNIZED_GLOBAL_OPS:
                result.issues.append(Issue("WARNING", sname, r, "G2",
                    f"Unrecognized global operation '{op_val}'. "
                    f"Recognized: {sorted(RECOGNIZED_GLOBAL_OPS)}"))


def _verify_joins_sheet(ws, server_sheets: list, result: VerificationResult):
    if ws.max_row is None or ws.max_row < 2:
        return

    max_col = ws.max_column or 0

    # Build header map
    hdr = {}
    for c in range(1, max_col + 1):
        n = _norm(ws.cell(1, c).value)
        if n:
            hdr[n] = c

    def _jget(row, name):
        c = hdr.get(name)
        return _clean(ws.cell(row, c).value) if c else ""

    aliases_per_sheet = {}

    for r in range(2, ws.max_row + 1):
        row_vals = [_clean(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        if all(_is_blank(v) for v in row_vals):
            continue

        sn = _jget(r, "sheet_name")
        join_alias = _jget(r, "join_alias")
        join_table = _jget(r, "join_table")
        join_type = _jget(r, "join_type")
        left_key = _jget(r, "left_key")
        right_key = _jget(r, "right_key")
        fetch_cols = _jget(r, "fetch_columns")

        # X7: Required fields
        for fname, fval in [("sheet_name", sn), ("join_alias", join_alias),
                            ("join_table", join_table), ("join_type", join_type),
                            ("left_key", left_key), ("right_key", right_key),
                            ("fetch_columns", fetch_cols)]:
            if _is_blank(fval):
                result.issues.append(Issue("ERROR", "_joins", r, "X7",
                    f"Required field '{fname}' is blank"))

        # X4: sheet_name matches server sheet
        if sn and sn not in server_sheets:
            result.issues.append(Issue("ERROR", "_joins", r, "X4",
                f"sheet_name '{sn}' does not match any server sheet. "
                f"Available: {server_sheets}"))

        # X5: join_type valid
        if join_type and join_type.upper() not in VALID_JOIN_TYPES:
            result.issues.append(Issue("ERROR", "_joins", r, "X5",
                f"join_type '{join_type}' is invalid. Must be: {VALID_JOIN_TYPES}"))

        # X6: alias unique per sheet
        if sn and join_alias:
            if sn not in aliases_per_sheet:
                aliases_per_sheet[sn] = {}
            if join_alias in aliases_per_sheet[sn]:
                prev = aliases_per_sheet[sn][join_alias]
                if join_table and prev != join_table.upper():
                    result.issues.append(Issue("ERROR", "_joins", r, "X6",
                        f"join_alias '{join_alias}' reused for different table "
                        f"'{join_table}' (previously '{prev}') in sheet '{sn}'"))
            else:
                aliases_per_sheet[sn][join_alias] = join_table.upper() if join_table else ""


def _verify_parameters_sheet(ws, result: VerificationResult):
    """P1, P2, P3: Validate the _parameters sheet."""
    if ws.max_row is None or ws.max_row < 2:
        return

    max_col = ws.max_column or 0

    # Build header map from row 1
    hdr = {}
    for c in range(1, max_col + 1):
        n = _norm(ws.cell(1, c).value)
        if n:
            hdr[n] = c

    # P1: Required headers
    missing_hdrs = PARAMETERS_REQUIRED_HEADERS - set(hdr.keys())
    if missing_hdrs:
        result.issues.append(Issue("ERROR", "_parameters", 1, "P1",
            f"Missing required headers: {missing_hdrs}"))
        return

    def _pget(row, name):
        c = hdr.get(name)
        return _clean(ws.cell(row, c).value) if c else ""

    seen_names = {}
    for r in range(2, ws.max_row + 1):
        row_vals = [_clean(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        if all(_is_blank(v) for v in row_vals):
            continue

        param_name = _pget(r, "parameter_name")

        # P2: parameter_name not blank
        if _is_blank(param_name):
            result.issues.append(Issue("ERROR", "_parameters", r, "P2",
                "parameter_name is blank"))
            continue

        # P3: No duplicate parameter_name
        param_upper = param_name.strip().upper()
        if param_upper in seen_names:
            result.issues.append(Issue("WARNING", "_parameters", r, "P3",
                f"Duplicate parameter_name '{param_name}' "
                f"(first seen at row {seen_names[param_upper]})"))
        else:
            seen_names[param_upper] = r


def _verify_cross_sheets(sheet_targets: dict, result: VerificationResult):
    names = list(sheet_targets.keys())
    ref_name = names[0]
    ref_cols = sheet_targets[ref_name]
    ref_set = set(ref_cols.keys())

    for other in names[1:]:
        other_cols = sheet_targets[other]
        other_set = set(other_cols.keys())

        # X1: Same column set
        missing = ref_set - other_set
        extra = other_set - ref_set
        if missing:
            result.issues.append(Issue("ERROR", other, None, "X1",
                f"Missing target columns vs '{ref_name}': {missing}"))
        if extra:
            result.issues.append(Issue("ERROR", other, None, "X1",
                f"Extra target columns vs '{ref_name}': {extra}"))

        # X2: Same dtypes
        for col in ref_set & other_set:
            if ref_cols[col]["dtype"].upper() != other_cols[col]["dtype"].upper():
                result.issues.append(Issue("ERROR", other, None, "X2",
                    f"Type mismatch for '{col}': '{ref_name}' has "
                    f"'{ref_cols[col]['dtype']}', '{other}' has '{other_cols[col]['dtype']}'"))

        # X3: Same PK markers
        for col in ref_set & other_set:
            if ref_cols[col]["pk"] != other_cols[col]["pk"]:
                result.issues.append(Issue("ERROR", other, None, "X3",
                    f"PK mismatch for '{col}': '{ref_name}' has "
                    f"'{ref_cols[col]['pk']}', '{other}' has '{other_cols[col]['pk']}'"))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    if len(sys.argv) < 2:
        print("Usage: python verify_excel_syntax.py <path_to_xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    res = verify(path)
    print(res.summary())
    sys.exit(0 if res.passed else 1)


if __name__ == "__main__":
    main()

