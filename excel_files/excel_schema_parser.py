"""
excel_schema_parser.py
-----------------------
Single source-of-truth Excel mapping parser for the ETL pipeline.

Reads Excel mapping specs using the project format (src_db, tgt_col_name headers)
and produces a structured dict with:
  - target_db / target_table (from filename)
  - unified target schema
  - per-server: source_tables, column_mappings, joins, global_transforms
  - cross-sheet consistency_errors
  - parameters (from _parameters sheet)

Call chain:
    custom_execution → generate_etl_runbook → excel_schema_parser → verify_excel_syntax

Usage (standalone — produces JSON):
    python excel_schema_parser.py

Programmatic (returns dict, no JSON written):
    from excel_schema_parser import parse_excel
    result = parse_excel("analytics_dw.public.dim_vehicle_master.xlsx")
"""

from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any

import openpyxl

# ═══════════════════════════════════════════════════════════════════════════
# Constants
# ═══════════════════════════════════════════════════════════════════════════

FILENAME_RE = re.compile(
    r'^(?P<target_db>[^.]+)\.(?P<target_schema>[^.]+)\.(?P<target_table>.+)\.xlsx$',
    re.IGNORECASE,
)

JOINS_SHEET = "_joins"

# Sheets to always skip (case-insensitive)
SKIP_SHEETS = {"_joins", "_parameters"}

# Global transforms marker and headers
GT_MARKER = "_global_transforms"
GT_HEADERS_EXPECTED = {"order", "operation", "parameters", "condition", "notes"}

# NULL-handling patterns in Notes column
NULL_PATTERNS = [
    (re.compile(r'DEFAULT\s+(\S+)\s+IF\s+NULL', re.I), r'fill:\1'),
    (re.compile(r'UNSPECIFIED\s+IF\s+NULL', re.I), 'fill:UNSPECIFIED'),
    (re.compile(r"DEFAULT\s+'([^']+)'", re.I), r'fill:\1'),
    (re.compile(r'DEFAULT\s+(\S+)', re.I), r'fill:\1'),
]

# Chained replace pattern
REPLACE_PATTERN = re.compile(
    r'replace\s*\(\s*"([^"]*)"\s*,\s*"([^"]*)"\s*\)', re.IGNORECASE,
)

# Joins required headers
JOINS_REQUIRED = {
    "sheet_name", "join_alias", "join_table",
    "join_type", "left_key", "right_key", "fetch_columns",
}


# ═══════════════════════════════════════════════════════════════════════════
# Data Models
# ═══════════════════════════════════════════════════════════════════════════

@dataclass
class TargetColumnDef:
    """A single column in the target table."""
    name: str
    dtype: str
    nullable: bool
    is_pk: bool
    pk_marker: str          # e.g. "Y", "N", ""
    default_val: str        # from tgt_default_val column
    force_dtype: str        # from force_dtype column
    null_handling: str      # explicit null_handling column or parsed from Notes


@dataclass
class SourceTableRef:
    """A distinct source table referenced in a server sheet."""
    database: str
    schema: str
    table: str
    columns: list[dict] = field(default_factory=list)   # [{name, dtype, nullable, is_pk}]


@dataclass
class ColumnMapping:
    """One row mapping: source field → target column with transform."""
    # Target side
    target_column: str
    target_dtype: str
    target_nullable: bool
    target_pk_marker: str
    target_default_val: str
    # Source side
    source_database: str
    source_schema: str
    source_table: str
    source_field: str
    source_dtype: str
    source_nullable: str
    source_is_pk: str
    source_is_unique: str
    # Transform
    transform_type: str     # direct, rename, cast, derived, constant, drop
    transform_rule: str     # the rule expression
    # Metadata
    force_dtype: str
    null_handling: str      # explicit column value or parsed from Notes
    notes: str


@dataclass
class JoinDef:
    """A join definition from the _joins sheet."""
    sheet_name: str
    join_alias: str
    join_table: str
    join_db: str
    join_schema: str
    join_type: str          # LEFT / INNER / RIGHT / FULL
    left_key: str           # first left key (convenience)
    right_key: str          # first right key (convenience)
    key_pairs: list[tuple[str, str]] = field(default_factory=list)
    fetch_columns: list[str] = field(default_factory=list)


@dataclass
class GlobalTransform:
    """One global transform rule."""
    order: int
    operation: str
    parameters: str
    condition: str
    notes: str


@dataclass
class ServerData:
    """All parsed data for one server sheet."""
    sheet_name: str
    source_tables: dict[str, SourceTableRef] = field(default_factory=dict)
    column_mappings: list[ColumnMapping] = field(default_factory=list)
    joins: list[JoinDef] = field(default_factory=list)
    global_transforms: list[GlobalTransform] = field(default_factory=list)
    target_columns: list[TargetColumnDef] = field(default_factory=list)


# ═══════════════════════════════════════════════════════════════════════════
# Helper Utilities
# ═══════════════════════════════════════════════════════════════════════════

def _clean(val: Any) -> str:
    """Return cleaned string from cell value, or empty string."""
    if val is None:
        return ""
    s = str(val).strip()
    s = s.replace('\xa0', ' ').strip()
    return s


def _norm(val: Any) -> str:
    """Normalised lowercase version of a cell value."""
    return _clean(val).lower()


def _bool_nullable(val: str) -> bool:
    """Parse Allow Nulls field → True if nullable."""
    v = val.strip().upper()
    return v not in ("N", "NO", "FALSE", "0")


def _parse_null_handling(notes: str) -> str:
    """Extract null-handling directive from Notes column."""
    if not notes:
        return ""
    for pattern, replacement in NULL_PATTERNS:
        m = pattern.search(notes)
        if m:
            if isinstance(replacement, str) and '\\' in replacement:
                return pattern.sub(replacement, m.group(0))
            return replacement if not m.groups() else f"fill:{m.group(1)}"
    return ""


def _source_table_key(db: str, schema: str, table: str) -> str:
    """Build a unique key for a source table."""
    parts = [p for p in [db, schema, table] if p]
    return ".".join(parts).upper()


# ═══════════════════════════════════════════════════════════════════════════
# Sheet Discovery
# ═══════════════════════════════════════════════════════════════════════════

def _should_skip_sheet(name: str) -> bool:
    """Check if a sheet should be skipped based on name."""
    nl = name.lower().strip()
    return nl in SKIP_SHEETS


def _is_server_sheet(ws) -> bool:
    """Check if worksheet has project-format headers (src_db + tgt_col_name)."""
    for row_idx in range(1, min(15, (ws.max_row or 0) + 1)):
        cells = []
        for col_idx in range(1, min((ws.max_column or 0) + 1, 25)):
            v = _norm(ws.cell(row_idx, col_idx).value)
            cells.append(v)
        if "src_db" in cells and "tgt_col_name" in cells:
            return True
    return False


def discover_sheets(wb) -> tuple[list[str], bool]:
    """
    Return (server_sheet_names, has_joins).

    Skips known non-data sheets and validates that remaining sheets
    have a recognisable header structure.
    """
    server_sheets = []
    has_joins = JOINS_SHEET in [s.lower() for s in wb.sheetnames]

    for name in wb.sheetnames:
        if _should_skip_sheet(name):
            continue
        ws = wb[name]
        if _is_server_sheet(ws):
            server_sheets.append(name)

    return server_sheets, has_joins


# ═══════════════════════════════════════════════════════════════════════════
# Header Detection (project format only)
# ═══════════════════════════════════════════════════════════════════════════

@dataclass
class HeaderMap:
    """Maps header names to column indices for one sheet."""
    header_row: int
    headers: dict[str, int]   # normalised name → col index


def detect_headers(ws) -> HeaderMap:
    """
    Scan rows 1–14 for the header row with project-format headers
    (src_db, tgt_col_name, etc.) and build a header→column index map.
    """
    header_row = -1
    max_col = min(ws.max_column or 1, 30)

    for row_idx in range(1, min(15, (ws.max_row or 1) + 1)):
        row_vals = [_norm(ws.cell(row_idx, c).value) for c in range(1, max_col + 1)]
        if "src_db" in row_vals and "tgt_col_name" in row_vals:
            header_row = row_idx
            break

    if header_row < 0:
        raise ValueError(f"Sheet '{ws.title}': cannot find header row with src_db and tgt_col_name")

    headers: dict[str, int] = {}
    for col_idx in range(1, max_col + 1):
        name = _norm(ws.cell(header_row, col_idx).value)
        if name:
            headers[name] = col_idx

    return HeaderMap(header_row=header_row, headers=headers)


# ═══════════════════════════════════════════════════════════════════════════
# Data Row Parsing
# ═══════════════════════════════════════════════════════════════════════════

def parse_data_rows(ws, hdr: HeaderMap) -> tuple[
    list[ColumnMapping], list[TargetColumnDef], dict[str, SourceTableRef]
]:
    """
    Parse all data rows from the header row to the global transforms marker or end.

    Returns (column_mappings, target_columns, source_tables).
    """
    mappings: list[ColumnMapping] = []
    target_cols: list[TargetColumnDef] = []
    source_tables: dict[str, SourceTableRef] = {}
    seen_tgt_cols: set[str] = set()

    def _get(row: int, name: str) -> str:
        col_idx = hdr.headers.get(name)
        if col_idx is None:
            return ""
        return _clean(ws.cell(row, col_idx).value)

    for row_idx in range(hdr.header_row + 1, (ws.max_row or hdr.header_row) + 1):
        # Check for global transforms marker
        cell_a = _clean(ws.cell(row_idx, 1).value)
        if cell_a.lower() == GT_MARKER:
            break

        # Check for legend row
        if cell_a.lower() == "transform type legend":
            break

        # Skip fully blank rows
        all_blank = True
        for c in range(1, min((ws.max_column or 1) + 1, 20)):
            if _clean(ws.cell(row_idx, c).value):
                all_blank = False
                break
        if all_blank:
            continue

        # Read target side
        tgt_col = _get(row_idx, "tgt_col_name")
        tgt_dtype = _get(row_idx, "tgt_dtype")
        tgt_nullable = _get(row_idx, "tgt_nullable")
        tgt_is_pk = _get(row_idx, "tgt_is_pk")
        tgt_default_val = _get(row_idx, "tgt_default_val")

        # Read source side
        src_db = _get(row_idx, "src_db")
        src_schema = _get(row_idx, "src_schema")
        src_table = _get(row_idx, "src_table")
        src_col = _get(row_idx, "src_col_name")
        src_dtype = _get(row_idx, "src_dtype")
        src_nullable = _get(row_idx, "src_nullable")
        src_is_pk = _get(row_idx, "src_is_pk")
        src_is_unique = _get(row_idx, "src_is_unique")

        # Read transform
        transform_type = _get(row_idx, "transform_type")
        transform_rule = _get(row_idx, "transform_rule")

        # Read metadata
        force_dtype = _get(row_idx, "force_dtype")
        null_handling_explicit = _get(row_idx, "null_handling")
        notes = _get(row_idx, "notes")

        if not tgt_col and not transform_type:
            continue
        if not tgt_col:
            continue

        # Determine null_handling: explicit column wins, else parse from Notes
        null_handling = null_handling_explicit or _parse_null_handling(notes)

        # Parse PK and nullable
        is_pk = tgt_is_pk.upper() in ("Y", "YES") if tgt_is_pk else False
        nullable = _bool_nullable(tgt_nullable) if tgt_nullable else True

        # Build target column def (deduplicated)
        tgt_col_upper = tgt_col.upper()
        if tgt_col_upper not in seen_tgt_cols:
            seen_tgt_cols.add(tgt_col_upper)
            target_cols.append(TargetColumnDef(
                name=tgt_col_upper,
                dtype=tgt_dtype,
                nullable=nullable,
                is_pk=is_pk,
                pk_marker=tgt_is_pk,
                default_val=tgt_default_val,
                force_dtype=force_dtype,
                null_handling=null_handling,
            ))

        # Build source table ref
        if src_db and src_table:
            st_key = _source_table_key(src_db, src_schema, src_table)
            if st_key not in source_tables:
                source_tables[st_key] = SourceTableRef(
                    database=src_db, schema=src_schema, table=src_table, columns=[],
                )
            existing_cols = {c["name"].upper() for c in source_tables[st_key].columns}
            if src_col and src_col.upper() not in existing_cols:
                source_tables[st_key].columns.append({
                    "name": src_col,
                    "dtype": src_dtype,
                    "nullable": src_nullable,
                    "is_pk": src_is_pk,
                })

        # Normalize transform_type
        tt = transform_type.lower().strip() if transform_type else "direct"

        # Build mapping
        mappings.append(ColumnMapping(
            target_column=tgt_col_upper,
            target_dtype=tgt_dtype,
            target_nullable=nullable,
            target_pk_marker=tgt_is_pk,
            target_default_val=tgt_default_val,
            source_database=src_db,
            source_schema=src_schema,
            source_table=src_table,
            source_field=src_col,
            source_dtype=src_dtype,
            source_nullable=src_nullable,
            source_is_pk=src_is_pk,
            source_is_unique=src_is_unique,
            transform_type=tt,
            transform_rule=transform_rule,
            force_dtype=force_dtype,
            null_handling=null_handling,
            notes=notes,
        ))

    return mappings, target_cols, source_tables


# ═══════════════════════════════════════════════════════════════════════════
# Global Transforms Parsing
# ═══════════════════════════════════════════════════════════════════════════

def parse_global_transforms(ws, hdr: HeaderMap) -> list[GlobalTransform]:
    """
    Find the _global_transforms marker row and parse the transform entries below it.
    """
    marker_row = -1
    for row_idx in range(hdr.header_row + 1, (ws.max_row or hdr.header_row) + 1):
        cell_a = _clean(ws.cell(row_idx, 1).value)
        if cell_a.lower() == GT_MARKER:
            marker_row = row_idx
            break

    if marker_row < 0:
        return []

    gt_hdr_row = marker_row + 1
    if gt_hdr_row > (ws.max_row or 0):
        return []

    gt_hdr_map: dict[str, int] = {}
    for col_idx in range(1, min((ws.max_column or 1) + 1, 15)):
        val = _norm(ws.cell(gt_hdr_row, col_idx).value)
        if val in GT_HEADERS_EXPECTED:
            gt_hdr_map[val] = col_idx

    if "order" not in gt_hdr_map or "operation" not in gt_hdr_map:
        return []

    transforms: list[GlobalTransform] = []
    for row_idx in range(gt_hdr_row + 1, (ws.max_row or gt_hdr_row) + 1):
        order_val = _clean(ws.cell(row_idx, gt_hdr_map["order"]).value)
        op_val = _clean(ws.cell(row_idx, gt_hdr_map["operation"]).value)

        if not order_val and not op_val:
            break

        try:
            order_int = int(order_val)
        except (ValueError, TypeError):
            break  # non-numeric order = end of section

        params = _clean(ws.cell(row_idx, gt_hdr_map.get("parameters", 0)).value) if "parameters" in gt_hdr_map else ""
        condition = _clean(ws.cell(row_idx, gt_hdr_map.get("condition", 0)).value) if "condition" in gt_hdr_map else ""
        notes = _clean(ws.cell(row_idx, gt_hdr_map.get("notes", 0)).value) if "notes" in gt_hdr_map else ""

        transforms.append(GlobalTransform(
            order=order_int,
            operation=op_val.upper(),
            parameters=params,
            condition=condition,
            notes=notes,
        ))

    return transforms


# ═══════════════════════════════════════════════════════════════════════════
# _joins Sheet Parsing
# ═══════════════════════════════════════════════════════════════════════════

def parse_joins_sheet(ws) -> dict[str, list[JoinDef]]:
    """
    Parse the _joins sheet into JoinDef objects grouped by sheet_name.
    Supports composite keys (multiple rows with same alias are merged).
    """
    header_map: dict[str, int] = {}
    for col_idx in range(1, (ws.max_column or 1) + 1):
        val = _norm(ws.cell(1, col_idx).value)
        if val:
            header_map[val] = col_idx

    # Temporary storage for composite key grouping
    grouped: dict[tuple, JoinDef] = {}
    order: list[tuple] = []

    for row_idx in range(2, (ws.max_row or 1) + 1):
        def _get(name: str) -> str:
            idx = header_map.get(name, 0)
            return _clean(ws.cell(row_idx, idx).value) if idx else ""

        sheet_name = _get("sheet_name")
        if not sheet_name:
            continue

        join_alias = _get("join_alias")
        join_table = _get("join_table")
        join_db = _get("join_db")
        join_schema = _get("join_schema")
        join_type = _get("join_type").upper() or "LEFT"
        left_key = _get("left_key")
        right_key = _get("right_key")
        fetch_raw = _get("fetch_columns")

        if not all([join_alias, join_table, left_key, right_key]):
            print(f"  WARNING: _joins row {row_idx} incomplete, skipping.")
            continue

        gkey = (sheet_name, join_alias, join_table)

        if gkey in grouped:
            # Composite key — add key pair
            grouped[gkey].key_pairs.append((left_key, right_key))
        else:
            fetch_cols = [c.strip() for c in fetch_raw.split(",") if c.strip()]
            jd = JoinDef(
                sheet_name=sheet_name,
                join_alias=join_alias,
                join_table=join_table,
                join_db=join_db,
                join_schema=join_schema,
                join_type=join_type,
                left_key=left_key,
                right_key=right_key,
                key_pairs=[(left_key, right_key)],
                fetch_columns=fetch_cols,
            )
            grouped[gkey] = jd
            order.append(gkey)

    # Group by sheet_name
    result: dict[str, list[JoinDef]] = {}
    for gkey in order:
        jd = grouped[gkey]
        result.setdefault(jd.sheet_name, []).append(jd)

    return result


# ═══════════════════════════════════════════════════════════════════════════
# _parameters Sheet Parsing
# ═══════════════════════════════════════════════════════════════════════════

def parse_parameters_sheet(wb) -> dict[str, str]:
    """
    Parse the _parameters sheet from a workbook.

    Returns a dict of {PARAMETER_NAME: parameter_value} for all non-blank rows.
    Returns empty dict if the sheet does not exist.
    """
    params_sheet_name = None
    for s in wb.sheetnames:
        if s.lower() == "_parameters":
            params_sheet_name = s
            break

    if params_sheet_name is None:
        return {}

    ws = wb[params_sheet_name]
    if (ws.max_row or 0) < 2:
        return {}

    hmap: dict[str, int] = {}
    for c in range(1, (ws.max_column or 1) + 1):
        name = _clean(ws.cell(1, c).value).lower()
        if name:
            hmap[name] = c

    name_col = hmap.get("parameter_name")
    value_col = hmap.get("parameter_value")
    if not name_col or not value_col:
        print(f"  WARNING: _parameters sheet missing required headers (parameter_name, parameter_value)")
        return {}

    parameters: dict[str, str] = {}
    for r in range(2, (ws.max_row or 1) + 1):
        param_name = _clean(ws.cell(r, name_col).value)
        param_value = _clean(ws.cell(r, value_col).value)
        if not param_name:
            continue
        parameters[param_name.strip().upper()] = param_value

    return parameters


# ═══════════════════════════════════════════════════════════════════════════
# Cross-Sheet Validation
# ═══════════════════════════════════════════════════════════════════════════

def validate_cross_sheet(servers: dict[str, ServerData]) -> list[str]:
    """
    Validate that all server sheets produce the same target table structure.

    Checks:
      - Same set of target column names (X1)
      - Same data types per column (X2)
      - Same PK markers per column (X3)
    """
    errors: list[str] = []
    sheet_names = list(servers.keys())
    if len(sheet_names) < 2:
        return errors

    ref_name = sheet_names[0]
    ref = servers[ref_name]
    ref_cols = {tc.name for tc in ref.target_columns}
    ref_types = {tc.name: tc.dtype for tc in ref.target_columns}
    ref_pks = {tc.name: tc.pk_marker for tc in ref.target_columns}

    for other_name in sheet_names[1:]:
        other = servers[other_name]
        other_cols = {tc.name for tc in other.target_columns}
        other_types = {tc.name: tc.dtype for tc in other.target_columns}
        other_pks = {tc.name: tc.pk_marker for tc in other.target_columns}

        missing_in_other = ref_cols - other_cols
        extra_in_other = other_cols - ref_cols
        if missing_in_other:
            errors.append(
                f"[X1] Sheet '{other_name}' missing target columns present in '{ref_name}': "
                f"{sorted(missing_in_other)}"
            )
        if extra_in_other:
            errors.append(
                f"[X1] Sheet '{other_name}' has extra target columns not in '{ref_name}': "
                f"{sorted(extra_in_other)}"
            )

        for col in ref_cols & other_cols:
            if (ref_types.get(col, "").upper() != other_types.get(col, "").upper()):
                errors.append(
                    f"[X2] Column '{col}' dtype mismatch: "
                    f"'{ref_name}'={ref_types[col]} vs '{other_name}'={other_types[col]}"
                )

        for col in ref_cols & other_cols:
            r_pk = (ref_pks.get(col) or "").upper()
            o_pk = (other_pks.get(col) or "").upper()
            if r_pk != o_pk:
                errors.append(
                    f"[X3] Column '{col}' PK marker mismatch: "
                    f"'{ref_name}'={r_pk!r} vs '{other_name}'={o_pk!r}"
                )

    return errors


# ═══════════════════════════════════════════════════════════════════════════
# Serialization
# ═══════════════════════════════════════════════════════════════════════════

def _serialize_server(sd: ServerData) -> dict:
    """Convert a ServerData to a JSON-serializable dict."""
    return {
        "source_tables": {
            key: {
                "database": st.database,
                "schema": st.schema,
                "table": st.table,
                "columns": st.columns,
            }
            for key, st in sd.source_tables.items()
        },
        "column_mappings": [asdict(m) for m in sd.column_mappings],
        "joins": [
            {
                "sheet_name": j.sheet_name,
                "join_alias": j.join_alias,
                "join_table": j.join_table,
                "join_db": j.join_db,
                "join_schema": j.join_schema,
                "join_type": j.join_type,
                "left_key": j.left_key,
                "right_key": j.right_key,
                "key_pairs": j.key_pairs,
                "fetch_columns": j.fetch_columns,
            }
            for j in sd.joins
        ],
        "global_transforms": [asdict(gt) for gt in sd.global_transforms],
    }


# ═══════════════════════════════════════════════════════════════════════════
# Top-Level Parser
# ═══════════════════════════════════════════════════════════════════════════

def parse_excel(
    excel_path: str | Path,
    run_syntax_check: bool = True,
) -> dict:
    """
    Main entry point. Parse an Excel mapping spec and return a dict.

    Parameters
    ----------
    excel_path : path to the .xlsx file
    run_syntax_check : if True, runs verify_excel_syntax.verify() first

    Returns
    -------
    dict with keys: file_name, target_db, target_schema, target_table,
                    target_columns, target_pks,
                    servers, consistency_errors, parameters
    """
    path = Path(excel_path)
    fname = path.name

    # ── Run syntax verification ──
    if run_syntax_check:
        from excel_files.verify_excel_syntax import verify as verify_syntax
        print(f"  Verifying syntax: {fname}")
        vresult = verify_syntax(str(path))
        if not vresult.passed:
            print(f"  {vresult.summary()}")
            raise ValueError(
                f"Excel syntax verification FAILED for '{fname}'. "
                f"Fix errors above before parsing."
            )
        errors = [i for i in vresult.issues if i.severity == "ERROR"]
        warnings = [i for i in vresult.issues if i.severity == "WARNING"]
        print(f"  Syntax OK ({len(errors)} errors, {len(warnings)} warnings)")
    else:
        print(f"  Skipping syntax verification (run_syntax_check=False)")

    # ── Extract target info from filename ──
    m = FILENAME_RE.match(fname)
    if m:
        target_db = m.group("target_db")
        target_schema = m.group("target_schema")
        target_table = f"{target_schema}.{m.group('target_table')}"
    else:
        parts = path.stem.split(".")
        if len(parts) >= 3:
            target_db = parts[0]
            target_schema = parts[1]
            target_table = ".".join(parts[1:])
        elif len(parts) == 2:
            target_db = parts[0]
            target_schema = ""
            target_table = parts[1]
        else:
            target_db = ""
            target_schema = ""
            target_table = path.stem

    wb = openpyxl.load_workbook(str(path), data_only=True)

    # ── Discover sheets ──
    server_sheet_names, has_joins = discover_sheets(wb)
    if not server_sheet_names:
        raise ValueError(f"No valid server sheets found in '{fname}'")

    print(f"  Server sheets: {server_sheet_names}")

    # ── Parse _joins sheet ──
    joins_by_sheet: dict[str, list[JoinDef]] = {}
    if has_joins:
        joins_ws_name = next(s for s in wb.sheetnames if s.lower() == JOINS_SHEET)
        print(f"  Found '{joins_ws_name}' sheet — parsing joins...")
        joins_by_sheet = parse_joins_sheet(wb[joins_ws_name])
        print(f"  Joins loaded for sheets: {list(joins_by_sheet.keys())}")

    # ── Parse _parameters sheet ──
    parameters = parse_parameters_sheet(wb)
    if parameters:
        print(f"  Parameters: {len(parameters)} parameter(s)")

    # ── Parse each server sheet ──
    servers: dict[str, ServerData] = {}

    for sheet_name in server_sheet_names:
        ws = wb[sheet_name]
        print(f"  Parsing sheet '{sheet_name}'...")

        hdr = detect_headers(ws)
        mappings, target_cols, source_tables = parse_data_rows(ws, hdr)
        print(f"    {len(mappings)} column mappings, {len(source_tables)} source table(s)")

        gt = parse_global_transforms(ws, hdr)
        if gt:
            print(f"    {len(gt)} global transform(s)")

        sheet_joins = joins_by_sheet.get(sheet_name, [])

        sd = ServerData(
            sheet_name=sheet_name,
            source_tables=source_tables,
            column_mappings=mappings,
            joins=sheet_joins,
            global_transforms=gt,
            target_columns=target_cols,
        )
        servers[sheet_name] = sd

    # ── Cross-sheet validation ──
    consistency_errors = validate_cross_sheet(servers)
    if consistency_errors:
        print(f"\n  ⚠ Cross-sheet consistency errors:")
        for err in consistency_errors:
            print(f"    {err}")

    # ── Build result ──
    first_server = servers[server_sheet_names[0]]
    pks = [tc.name for tc in first_server.target_columns if tc.is_pk]

    result = {
        "file_name": fname,
        "target_db": target_db,
        "target_schema": target_schema,
        "target_table": target_table,
        "target_columns": [
            {
                "name": tc.name,
                "dtype": tc.dtype,
                "nullable": tc.nullable,
                "is_pk": tc.is_pk,
                "pk_marker": tc.pk_marker,
                "default_val": tc.default_val,
                "force_dtype": tc.force_dtype,
                "null_handling": tc.null_handling,
            }
            for tc in first_server.target_columns
        ],
        "target_pks": pks,
        "servers": {
            name: _serialize_server(sd) for name, sd in servers.items()
        },
        "consistency_errors": consistency_errors,
        "parameters": parameters,
    }

    wb.close()
    return result


# ═══════════════════════════════════════════════════════════════════════════
# Configuration — edit these variables before running standalone
# ═══════════════════════════════════════════════════════════════════════════

# EXCEL_FILE = "analytics_dw.public.dim_vehicle_master.xlsx"
# EXCEL_FILE = "analytics_dw.public.fact_commercial.xlsx"
EXCEL_FILE = "analytics_dw.public.fact_production.xlsx"

# Output JSON path (None = auto-generate as <stem>_parsed.json in same folder)
OUTPUT_JSON = None

# Pretty-print the JSON output
PRETTY_PRINT = True

# Run syntax verification before parsing
RUN_SYNTAX_CHECK = True


if __name__ == "__main__":
    excel_path = Path(EXCEL_FILE)
    if not excel_path.exists():
        print(f"ERROR: {excel_path} not found.", file=sys.stderr)
        sys.exit(1)

    print(f"\nParsing: {excel_path.name}")
    try:
        result = parse_excel(excel_path, run_syntax_check=RUN_SYNTAX_CHECK)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)

    out_path = Path(OUTPUT_JSON) if OUTPUT_JSON else excel_path.with_name(excel_path.stem + "_parsed.json")
    out_path.write_text(
        json.dumps(result, indent=2 if PRETTY_PRINT else None, default=str),
        encoding="utf-8",
    )

    n_servers = len(result["servers"])
    n_errors = len(result["consistency_errors"])
    print(f"\n  OK  {n_servers} server(s) parsed → {out_path}")
    if n_errors:
        print(f"  ⚠  {n_errors} consistency error(s) — see output JSON")


